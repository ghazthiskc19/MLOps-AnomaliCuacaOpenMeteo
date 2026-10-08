"""
scripts/generate_lk05_report.py - Generator Laporan Akademik LK-05 (DOCX & PDF).
Laporan Manajemen Data Versioning Menggunakan DVC. Format (font, tabel, margin, kotak kode, penomoran bab)
mengikuti LK-04 dengan menggunakan ulang helper dari generate_lk04_report.py, ditambah Daftar Isi otomatis
(field TOC Word dengan nomor halaman) dan nomor halaman pada footer.
"""

import os
from pathlib import Path
import subprocess
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from generate_lk04_report import (  # noqa: E402
    REFERENCE_DOCX_PATH,
    REPO_ROOT,
    add_callout,
    add_custom_table,
    add_figure,
    add_heading_1,
    add_heading_2,
    add_p,
    format_run,
)

LK05_DIR = REPO_ROOT / "Documentation" / "LK05"
BASENAME = "LK05-Muhammad Ghazy Humaidi-245150200111071"
DOCX_OUT_PATH = LK05_DIR / f"{BASENAME}.docx"
PDF_OUT_PATH = LK05_DIR / f"{BASENAME}.pdf"
PUSH_EVIDENCE_PATH = LK05_DIR / "dvc_push_evidence.txt"  # opsional: keluaran `dvc push` dari jaringan UB


ASSETS_DIR = REPO_ROOT / "Documentation" / "assets"
LOGO_FILKOM = ASSETS_DIR / "logo-filkom.png"  # diekstrak dari cover LK-03
LOGO_UB = ASSETS_DIR / "logo-ub.png"


def add_logo(doc, image_path, width_in, space_after=10):
    """Menambahkan logo di tengah halaman sampul (format cover LK-03)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.add_run().add_picture(str(image_path), width=Inches(width_in))


def _set_outline_level(paragraph, level):
    """Menandai paragraf sebagai entri Daftar Isi (outline level) tanpa mengubah tampilan LK-04."""
    paragraph._p.get_or_add_pPr().append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{level}"/>'))


def h1(doc, text):
    p = add_heading_1(doc, text)
    _set_outline_level(p, 0)
    return p


def h2(doc, text):
    p = add_heading_2(doc, text)
    _set_outline_level(p, 1)
    return p


def _add_field(paragraph, instruction, placeholder, size_pt=12, bold=False):
    """Menyisipkan field kompleks Word (TOC / PAGE) yang diperbarui otomatis oleh Word."""
    def fld(kind):
        r = paragraph.add_run()
        format_run(r, size_pt=size_pt, bold=bold)
        r._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="{kind}"/>'))

    fld("begin")
    r = paragraph.add_run()
    format_run(r, size_pt=size_pt, bold=bold)
    r._r.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {instruction} </w:instrText>'))
    fld("separate")
    r = paragraph.add_run(placeholder)
    format_run(r, size_pt=size_pt, bold=bold)
    fld("end")


def add_page_number_footer(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_field(p, "PAGE", "1", size_pt=10)


def build_lk05_docx():
    print("Membangun dokumen LK-05 DOCX...")
    if REFERENCE_DOCX_PATH.exists():
        doc = docx.Document(str(REFERENCE_DOCX_PATH))
        body = doc._body._element
        for c in [c for c in body if not c.tag.endswith("sectPr")]:
            body.remove(c)
    else:
        doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Pt(65)
        section.bottom_margin = Pt(65)
        section.left_margin = Pt(72)
        section.right_margin = Pt(72)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
    add_page_number_footer(doc)

    C = WD_ALIGN_PARAGRAPH.CENTER

    # ------------------------------------------------------------------ COVER
    add_logo(doc, LOGO_FILKOM, width_in=2.9, space_after=14)
    add_p(doc, "LEMBAR KERJA-05 (LK-05)", align=C, size_pt=20, bold=True, space_after=10)
    add_p(doc, "Manajemen Data Versioning Menggunakan DVC", align=C, size_pt=17, bold=True, space_after=8)
    add_p(
        doc,
        "Sistem Otomatis Deteksi Anomali Data Cuaca Open-Meteo untuk Smart Farming Berbasis MLOps",
        align=C, size_pt=13.5, bold=True, space_after=12,
    )
    add_p(doc, "MLOPS TIF-B 2026", align=C, size_pt=15, bold=True, space_after=26)
    add_p(doc, "Nama Mahasiswa:", align=C, bold=True, space_after=2)
    add_p(doc, "Muhammad Ghazy Humaidi (245150200111071)", align=C, space_after=14)
    add_p(doc, "Dosen Pengampu:", align=C, bold=True, space_after=2)
    add_p(doc, "Rizal Setya Perdana, S.Kom., M.Kom., Ph.D.", align=C, space_after=16)
    add_logo(doc, LOGO_UB, width_in=2.0, space_after=14)
    add_p(doc, "Program Studi Teknik Informatika", align=C, bold=True, space_after=2)
    add_p(doc, "Jurusan Teknik Informatika", align=C, bold=True, space_after=2)
    add_p(doc, "Fakultas Ilmu Komputer, Universitas Brawijaya", align=C, bold=True, space_after=10)
    add_p(doc, "2026", align=C, bold=True, space_after=0)
    doc.add_page_break()

    # ------------------------------------------------------- RINGKASAN & TOC
    add_heading_1(doc, "RINGKASAN DAN DAFTAR ISI")
    add_custom_table(doc, ["Informasi Akademik", "Detail Implementasi Mahasiswa"], [
        ["Nama Mahasiswa / NIM", "Muhammad Ghazy Humaidi / 245150200111071"],
        ["Mata Kuliah / Dosen", "MLOps — TIF-B 2026 / Rizal Setya Perdana, S.Kom., M.Kom., Ph.D."],
        ["Institusi", "Teknik Informatika, FILKOM, Universitas Brawijaya"],
        ["Repositori GitHub & Branch", "MLOps-AnomaliCuacaOpenMeteo (feat/lk05-dvc-versioning)"],
        ["Cakupan Tugas LK-05", "Inisialisasi DVC, pelacakan dataset, simulasi data baru, audit diff antar versi"],
        ["Hasil Utama", "2 versi dataset terlacak: dataset-v1.0.0 (720 baris) → dataset-v2.0.0 (936 baris)"],
    ], col_widths=[2.3, 4.4])

    add_p(doc, "DAFTAR ISI", bold=True, size_pt=12.5, space_before=6, space_after=4)
    toc_p = add_p(doc, "", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)
    _add_field(toc_p, 'TOC \\o "1-2" \\h \\z \\u', "Daftar isi diperbarui otomatis oleh Microsoft Word.", size_pt=11)
    doc.add_page_break()

    # ------------------------------------------------------------------ BAB 1
    h1(doc, "BAB 1: PENDAHULUAN DAN KONSEP DATA VERSIONING")
    h2(doc, "1.1 Latar Belakang dan Tujuan")
    add_p(doc, "Pada LK-04, pipeline ingestion Apache Airflow menarik data cuaca Open-Meteo setiap jam dan menyimpannya pada rolling buffer 30 hari (720 baris) di data/raw/weather_raw_current.csv. Karena buffer bersifat bergulir, isi berkas yang sama berubah setiap jam: data tertua keluar, data terbaru masuk. Git pada dasarnya tidak dirancang untuk menyimpan berkas data yang berubah terus-menerus dan berukuran besar atau biner, sehingga riwayat repositori akan membengkak dan isi data tertentu tidak dapat dipanggil kembali secara bermakna.")
    add_p(doc, "LK-05 bertujuan mengintegrasikan DVC (Data Version Control) ke dalam repositori agar: (1) dataset awal hasil LK-04 dapat dilacak versinya tanpa membebani Git, (2) penambahan data baru (simulasi continual learning) tercatat sebagai versi baru dengan silsilah yang jelas, dan (3) dua versi dataset dapat dibandingkan secara objektif.")

    h2(doc, "1.2 Konsep Dasar DVC pada Proyek Ini")
    add_p(doc, "DVC memisahkan satu versi data menjadi dua bagian yang disimpan di tempat berbeda. Pemisahan inilah yang menjaga repositori Git tetap ringan:")
    add_custom_table(doc, ["Komponen", "Lokasi Penyimpanan", "Fungsi"], [
        ["File penunjuk (.dvc)", "Git (GitHub)", "Teks kecil berisi hash MD5, ukuran, dan path berkas. Menjadi identitas satu versi data."],
        ["Cache DVC (.dvc/cache)", "Lokal (di-ignore Git)", "Salinan isi data per hash. Memungkinkan berpindah versi tanpa unduhan ulang."],
        ["Remote storage", "Folder ~/dvc-store pada VPS (SSH)", "Penyimpanan eksternal di luar repositori untuk berbagi dan mencadangkan isi data."],
        ["Tag Git (dataset-vX.Y.Z)", "Git (GitHub)", "Label yang mudah dibaca untuk menunjuk commit dengan versi data tertentu."],
    ], col_widths=[1.7, 1.9, 3.1])

    # ------------------------------------------------------------------ BAB 2
    h1(doc, "BAB 2: SETUP DVC DAN REMOTE STORAGE")
    h2(doc, "2.1 Instalasi dan Inisialisasi")
    add_p(doc, "DVC dipasang dengan ekstensi SSH karena remote storage menggunakan protokol SSH, dan versinya dikunci pada requirements.txt agar lingkungan dapat direproduksi. Pekerjaan dilakukan pada branch terpisah sesuai GitHub Flow (feat/lk05-dvc-versioning).")
    add_callout(doc, [
        "$ git checkout -b feat/lk05-dvc-versioning",
        "$ pip install \"dvc[ssh]==3.67.1\"      # dicatat pada requirements.txt",
        "$ dvc init",
        "$ git commit -m \"chore(dvc): initialize DVC\"",
    ], title="Inisialisasi DVC")
    add_p(doc, "Perintah dvc init membentuk direktori .dvc/ (konfigurasi dan cache), berkas .dvc/.gitignore, serta .dvcignore. Berkas-berkas konfigurasi ini yang di-commit ke Git, sedangkan cache lokal tidak.")

    h2(doc, "2.2 Konfigurasi Remote Storage di Luar Repositori")
    add_p(doc, "Sebagai penyimpanan eksternal (butir opsional tugas), digunakan folder ~/dvc-store pada VPS Ubuntu yang sama dengan server Airflow, diakses melalui SSH. Pendekatan ini tidak memerlukan instalasi DVC pada VPS karena VPS hanya berperan sebagai gudang berkas.")
    add_callout(doc, [
        "$ ssh member@10.34.211.189 \"mkdir -p ~/dvc-store\"",
        "$ dvc remote add -d vps ssh://member@10.34.211.189/home/member/dvc-store",
        "$ dvc remote list",
        "vps     ssh://member@10.34.211.189/home/member/dvc-store        (default)",
        "$ git add .dvc/config && git commit -m \"chore(dvc): add SSH remote on VPS\"",
    ], title="Konfigurasi remote DVC")
    add_p(doc, "Catatan operasional: VPS hanya dapat diakses dari jaringan UB, sehingga dvc push dan dvc pull hanya berfungsi saat terhubung ke jaringan tersebut. Untuk penggunaan lintas jaringan, remote dapat dialihkan ke S3, GCS, atau MinIO tanpa mengubah alur kerja yang dijelaskan pada laporan ini.")

    # ------------------------------------------------------------------ BAB 3
    h1(doc, "BAB 3: PELACAKAN DATASET DAN SIMULASI CONTINUAL LEARNING")
    h2(doc, "3.1 Versi 1: Pelacakan Dataset Awal (Baseline LK-04)")
    add_p(doc, "Dataset awal adalah buffer mentah hasil LK-04 berisi 720 baris (2026-08-29 s.d. 2026-09-27, 34.671 byte). Berkas ini dilacak dengan dvc add, yang menghitung hash MD5 isi berkas dan menulis file penunjuk yang kemudian di-commit serta diberi tag.")
    add_callout(doc, [
        "$ dvc add data/raw/weather_raw_current.csv",
        "$ git add data/raw/weather_raw_current.csv.dvc",
        "$ git commit -m \"data: track raw weather dataset v1 (720 rows, 2026-08-29 to 2026-09-27)\"",
        "$ git tag dataset-v1.0.0",
        "",
        "# isi data/raw/weather_raw_current.csv.dvc (v1)",
        "outs:",
        "- md5: e42dc025e5df6d31f780db84ee5836d3",
        "  size: 34671",
        "  hash: md5",
        "  path: weather_raw_current.csv",
    ], title="Versi 1: dvc add dan file penunjuk")
    add_p(doc, "Berkas CSV itu sendiri tidak masuk ke Git: pola data/raw/*.csv sudah terdaftar pada .gitignore proyek, sehingga yang tercatat di Git hanya file penunjuk .dvc berukuran puluhan byte.")

    h2(doc, "3.2 Simulasi Continual Learning: Pengambilan Data Tambahan")
    add_p(doc, "Skrip ingestion dijalankan kembali untuk mengambil data terbaru dari Open-Meteo. Ukuran buffer default dikunci 720 baris sehingga data baru hanya menggeser jendela (baris lama keluar) dan jumlah baris tidak bertambah. Agar penambahan baris dapat didemonstrasikan sesuai tugas, buffer diperbesar menjadi 1000 baris melalui parameter --buffer-size. Skrip menggabungkan berkas lokal dengan hasil fetch 30 hari terakhir, melakukan deduplikasi timestamp (keep='last'), dan menyimpan secara atomik.")
    add_callout(doc, [
        "$ python src/ingest_data.py --buffer-size 1000",
        "# hasil: 936 baris data (937 baris dengan header), 2026-08-29 s.d. 2026-10-06 23:00",
    ], title="Simulasi data baru")
    add_p(doc, "Sumber data live tetap berasal dari pipeline Airflow per jam pada VPS (LK-04). Pengambilan berkala, toleransi galat jaringan (retry dengan exponential backoff, fail-fast pada HTTP 4xx), dan penulisan atomik pada skrip ingestion memastikan data dinamis dapat ditarik otomatis dan konsisten.")

    h2(doc, "3.3 Versi 2: Pelacakan Dataset yang Diperbarui")
    add_callout(doc, [
        "$ dvc add data/raw/weather_raw_current.csv",
        "$ git add data/raw/weather_raw_current.csv.dvc",
        "$ git commit -m \"data: update raw weather dataset v2 (new rows from ingestion)\"",
        "$ git tag dataset-v2.0.0",
        "",
        "# isi data/raw/weather_raw_current.csv.dvc (v2)",
        "outs:",
        "- md5: 8e4412625620af93ac368a5692e1a399",
        "  size: 45033",
        "  hash: md5",
        "  path: weather_raw_current.csv",
    ], title="Versi 2: hash pada file penunjuk berubah")
    add_p(doc, "Ringkasan perbandingan karakteristik kedua versi dataset:")
    add_custom_table(doc, ["Atribut", "dataset-v1.0.0", "dataset-v2.0.0"], [
        ["Jumlah baris data", "720", "936 (+216 baris)"],
        ["Rentang waktu", "2026-08-29 00:00 s.d. 2026-09-27 23:00", "2026-08-29 00:00 s.d. 2026-10-06 23:00"],
        ["Ukuran berkas", "34.671 byte", "45.033 byte"],
        ["MD5", "e42dc025e5df6d31f780db84ee5836d3", "8e4412625620af93ac368a5692e1a399"],
        ["Nilai kosong (NaN)", "0", "0"],
        ["Rata-rata suhu / kelembaban", "23,35 °C / 72,19 %", "23,54 °C / 73,05 %"],
    ], col_widths=[1.9, 2.4, 2.4])

    # ------------------------------------------------------------------ BAB 4
    h1(doc, "BAB 4: AUDIT DAN PERBANDINGAN VERSI DATASET")
    h2(doc, "4.1 Status dan Diff Metadata dengan DVC")
    add_p(doc, "dvc status memverifikasi bahwa isi berkas di workspace sesuai dengan hash pada file penunjuk, sedangkan dvc diff membandingkan dua versi data berdasarkan tag Git. Hasilnya menunjukkan berkas yang sama berstatus Modified dengan hash berubah dari e42dc025 menjadi 8e441262.")
    add_callout(doc, [
        "$ dvc status",
        "Data and pipelines are up to date.",
        "",
        "$ dvc diff dataset-v1.0.0 dataset-v2.0.0",
        "Modified:",
        "    data\\raw\\weather_raw_current.csv",
        "",
        "files summary: 1 modified",
        "",
        "$ dvc diff --show-hash dataset-v1.0.0 dataset-v2.0.0",
        "Modified:",
        "    e42dc025..8e441262  data\\raw\\weather_raw_current.csv",
        "",
        "files summary: 1 modified",
    ], title="Audit versi dengan DVC")

    h2(doc, "4.2 Perubahan File Penunjuk pada Git")
    add_p(doc, "Silsilah data juga terekam pada riwayat Git. Perbandingan file penunjuk antar tag memperlihatkan hanya dua nilai yang berubah, yaitu hash dan ukuran, sementara data berukuran puluhan kilobyte tidak ikut tersimpan di Git.")
    add_callout(doc, [
        "$ git diff dataset-v1.0.0 dataset-v2.0.0 -- data/raw/weather_raw_current.csv.dvc",
        " outs:",
        "-- md5: e42dc025e5df6d31f780db84ee5836d3",
        "-  size: 34671",
        "+- md5: 8e4412625620af93ac368a5692e1a399",
        "+  size: 45033",
        "   hash: md5",
        "   path: weather_raw_current.csv",
    ], title="Diff file penunjuk .dvc antar versi")

    h2(doc, "4.3 Analisis Isi Perbedaan Data")
    add_p(doc, "Perbandingan isi kedua versi (berdasarkan timestamp) memberikan temuan yang penting bagi reproduksibilitas:")
    add_p(doc, "1. Penambahan data: 216 baris baru (9 hari x 24 jam, 2026-09-28 s.d. 2026-10-06) sehingga total menjadi 936 baris.")
    add_p(doc, "2. Revisi data lama: dari 720 baris yang tumpang tindih, 17 baris (seluruhnya pada 2026-09-27 pukul 07:00 s.d. 23:00) memiliki nilai yang berbeda pada v2: suhu 17 baris, kelembaban 16, curah hujan 6, kelembaban tanah 16, radiasi 10, dan kecepatan angin 16. Jam-jam tersebut berada di ujung rentang v1 saat data diambil, kemungkinan besar nilai prakiraan yang kemudian direvisi oleh model Open-Meteo. Dengan demikian, berkas dengan nama yang sama dapat berisi nilai berbeda untuk jam yang sama.")
    add_p(doc, "3. Dampak pada label ground truth: pipeline preprocessing yang sama (tanpa perubahan kode) menghasilkan distribusi kelas berbeda pada kedua versi, seperti ditunjukkan tabel berikut.")
    add_custom_table(doc, ["Kelas Anomali", "Hasil pada v1 (720 baris)", "Hasil pada v2 (936 baris)"], [
        ["Class 0: Normal State", "666 (92,50%)", "867 (92,63%)"],
        ["Class 1: Hardware Fault", "0 (0,00%)", "0 (0,00%)"],
        ["Class 2: Extreme Environmental", "54 (7,50%)", "69 (7,37%)"],
    ], col_widths=[2.4, 2.15, 2.15])

    h2(doc, "4.4 Pemulihan Versi Tertentu (Checkout)")
    add_p(doc, "Karena setiap versi terikat pada tag Git dan hash DVC, dataset dapat dikembalikan ke kondisi persis saat versi tersebut dibuat. Pengujian berikut menunjukkan berkas data/raw/weather_raw_current.csv berpindah antara 720 baris (v1) dan 936 baris (v2) hanya dengan berpindah tag dan menjalankan dvc checkout.")
    add_callout(doc, [
        "$ git checkout dataset-v1.0.0 && dvc checkout",
        "M       data\\raw\\weather_raw_current.csv",
        "$ wc -l data/raw/weather_raw_current.csv",
        "721 data/raw/weather_raw_current.csv           # 720 baris data + 1 header",
        "$ tail -1 data/raw/weather_raw_current.csv",
        "2026-09-27 23:00:00,21.4,92,0.0,0.151,0.0,1.4",
        "",
        "$ git checkout feat/lk05-dvc-versioning && dvc checkout",
        "M       data\\raw\\weather_raw_current.csv",
        "$ wc -l data/raw/weather_raw_current.csv",
        "937 data/raw/weather_raw_current.csv           # 936 baris data + 1 header",
        "$ tail -1 data/raw/weather_raw_current.csv",
        "2026-10-06 23:00:00,21.5,95,0.0,0.213,0.0,5.8",
        "",
        "$ dvc status",
        "Data and pipelines are up to date.",
    ], title="Berpindah antar versi dataset")

    if PUSH_EVIDENCE_PATH.exists():
        h2(doc, "4.5 Sinkronisasi ke Remote Storage")
        add_p(doc, "Isi data seluruh versi dikirim ke remote storage dengan dvc push sebelum file penunjuk dikirim ke GitHub, agar penunjuk di Git tidak pernah mengacu pada data yang belum tersedia di remote.")
        add_callout(doc, PUSH_EVIDENCE_PATH.read_text(encoding="utf-8").strip().splitlines(), title="Keluaran dvc push")

    # ------------------------------------------------------------------ BAB 5
    h1(doc, "BAB 5: PENTINGNYA DATA VERSIONING UNTUK REPRODUKSIBILITAS MODEL")
    h2(doc, "5.1 Mengapa Versioning Dataset Krusial")
    add_p(doc, "Hasil sebuah model ditentukan oleh tiga hal: kode, konfigurasi, dan data. Git telah menjamin dua yang pertama, tetapi tanpa versioning data, eksperimen tidak dapat direproduksi. Bukti nyata muncul pada laporan ini: dua versi dataset menghasilkan jumlah baris, nilai pada jam yang sama, dan distribusi label yang berbeda (Class 2 berubah dari 54 menjadi 69 baris) walaupun kode preprocessing identik. Model yang dilatih pada v1 dan v2 akan berbeda, dan tanpa tag data, tidak ada cara untuk membuktikan model mana dilatih dengan data apa.")
    add_p(doc, "Pada proyek ini masalahnya lebih tajam karena buffer 30 hari bergulir: data lama otomatis terhapus dari berkas sumber dan nilai prakiraan dapat direvisi. Tanpa snapshot, kondisi data pada saat training hari ini tidak akan pernah dapat dipulihkan minggu depan. DVC menjadikan setiap snapshot yang dipilih permanen, terverifikasi hash, dan dapat dipulihkan kapan saja.")

    h2(doc, "5.2 Alur Reproduksi dan Kapan Versi Dibuat")
    add_p(doc, "Satu versi data = satu commit Git (kode + file penunjuk) = satu tag. Untuk mereproduksi eksperimen, cukup memulihkan tag tersebut:")
    add_callout(doc, [
        "git checkout dataset-v1.0.0      # kode + file penunjuk .dvc versi 1",
        "dvc checkout                     # (atau dvc pull pada clone baru) isi data persis versi 1",
        "python src/preprocess.py         # preprocessing menghasilkan fitur & label yang sama",
    ], title="Reproduksi eksperimen")
    add_p(doc, "Versi dataset tidak dibuat setiap jam. Pengambilan per jam adalah tugas Airflow dan hanya memperbarui data hidup. Versi dibuat secara sengaja pada momen yang membutuhkan reproduksibilitas: baseline awal, sebelum training model, dan saat retraining terjadwal atau dipicu drift. Pada tahap training berikutnya, tag data (misalnya dataset-v2.0.0) dicatat bersama metrik pada experiment tracking sehingga setiap model memiliki asal data yang jelas.")

    h2(doc, "5.3 Alur Penambahan Versi Data dan Keterbatasan")
    add_custom_table(doc, ["Langkah", "Perintah", "Tujuan"], [
        ["1. Ambil data baru", "python src/ingest_data.py --buffer-size 1000", "Menambah baris pada dataset"],
        ["2. Catat versi baru", "dvc add data/raw/weather_raw_current.csv", "Menghitung hash dan memperbarui file .dvc"],
        ["3. Commit dan tag", "git add *.dvc; git commit; git tag dataset-vX.Y.Z", "Mengikat versi data ke riwayat Git"],
        ["4. Kirim data", "dvc push --all-tags", "Mengirim isi data ke remote (dahulukan)"],
        ["5. Kirim penunjuk", "git push origin <branch> --tags", "Membagikan file penunjuk via GitHub"],
        ["6. Audit", "dvc status; dvc diff <tagA> <tagB>", "Memastikan silsilah dan perbedaan versi"],
    ], col_widths=[1.5, 3.0, 2.2])
    add_p(doc, "Keterbatasan: (1) remote pada VPS hanya dapat diakses dari jaringan UB; (2) pada LK-05 yang dilacak hanya data mentah (data/raw), sedangkan data/processed adalah turunan deterministik dari data mentah dan kode preprocessing sehingga dapat dibangun ulang; (3) pembuatan versi dilakukan manual pada momen penting, bukan otomatis per jam.")

    # ------------------------------------------------------------------ BAB 6
    h1(doc, "BAB 6: KESIMPULAN")
    add_p(doc, "Seluruh sasaran LK-05 telah terpenuhi:")
    add_p(doc, "1. DVC berhasil diinisialisasi dan konfigurasinya di-commit ke Git, lengkap dengan remote storage eksternal berbasis SSH.")
    add_p(doc, "2. Dataset awal hasil LK-04 (720 baris) dilacak sebagai dataset-v1.0.0, dan penambahan data melalui skrip ingestion menghasilkan dataset-v2.0.0 (936 baris) dengan hash yang berubah.")
    add_p(doc, "3. Audit melalui dvc status, dvc diff, dan git diff membuktikan transisi versi (Modified: e42dc025 menjadi 8e441262), dan checkout membuktikan kedua versi dapat dipulihkan secara persis.")
    add_p(doc, "4. Analisis isi menunjukkan perbedaan nyata antar versi (216 baris baru, 17 baris direvisi, distribusi label berubah), yang menegaskan bahwa versioning dataset adalah syarat reproduksibilitas model.")
    add_custom_table(doc, ["Luaran yang Diminta", "Lokasi pada Repositori"], [
        ["File penunjuk .dvc dan .gitignore", "data/raw/weather_raw_current.csv.dvc; .dvc/.gitignore; .gitignore"],
        ["Laporan singkat (PDF) dengan dvc status/diff", "Documentation/LK05/" + BASENAME + ".pdf"],
        ["Dokumentasi README alur versi data", "README.md, bagian \"Data Versioning (DVC)\""],
    ], col_widths=[2.6, 4.1])

    DOCX_OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DOCX_OUT_PATH))
    print(f"Berhasil membuat berkas DOCX di: {DOCX_OUT_PATH}")


def update_toc_and_export_pdf():
    """Membuka DOCX di Word: perbarui Daftar Isi & nomor halaman, simpan, lalu ekspor PDF."""
    print("Memperbarui Daftar Isi dan mengekspor PDF via Word COM...")
    vbs = REPO_ROOT / "scripts" / "_export_lk05.vbs"
    d = str(DOCX_OUT_PATH).replace("\\", "/")
    p = str(PDF_OUT_PATH).replace("\\", "/")
    vbs.write_text(f'''On Error Resume Next
Set word = CreateObject("Word.Application")
If Err.Number <> 0 Then
    WScript.Echo "Error: Word COM tidak tersedia. " & Err.Description
    WScript.Quit 1
End If
word.Visible = False
word.DisplayAlerts = 0
Set fso = CreateObject("Scripting.FileSystemObject")
inPath = fso.GetAbsolutePathName("{d}")
outPath = fso.GetAbsolutePathName("{p}")
Set doc = word.Documents.Open(inPath)
If Err.Number <> 0 Then
    WScript.Echo "Error: Gagal membuka dokumen. " & Err.Description
    word.Quit
    WScript.Quit 1
End If
doc.Repaginate
doc.TablesOfContents(1).Update
doc.Repaginate
doc.TablesOfContents(1).Update
doc.Fields.Update
doc.Save
doc.ExportAsFixedFormat outPath, 17, False, 0, 0, 1, 1, 0, True, True, 0, True, True, False
If Err.Number <> 0 Then
    WScript.Echo "Error: Gagal ekspor PDF. " & Err.Description
    doc.Close False
    word.Quit
    WScript.Quit 1
End If
doc.Close False
word.Quit
WScript.Echo "OK"
''', encoding="utf-8")
    try:
        res = subprocess.run(["cscript", "//nologo", str(vbs)], capture_output=True, text=True, check=True)
        print(res.stdout.strip())
        if PDF_OUT_PATH.exists():
            print(f"PDF: {PDF_OUT_PATH} ({PDF_OUT_PATH.stat().st_size / 1024:.1f} KB)")
    finally:
        if vbs.exists():
            os.remove(vbs)


def main():
    build_lk05_docx()
    update_toc_and_export_pdf()


if __name__ == "__main__":
    main()
